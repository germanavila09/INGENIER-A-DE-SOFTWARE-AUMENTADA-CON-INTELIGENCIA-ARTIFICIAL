"""Extracción de texto de los documentos del bucket, dividida en secciones.

Una sección es la unidad que se cita: la página de un PDF, la diapositiva de un
PowerPoint, la hoja de un Excel o una parte de ~4 000 caracteres en Word y texto plano.
Las librerías de cada formato se importan solo cuando se necesitan.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import PurePosixPath

TAM_PARTE = 4000          # caracteres por sección en formatos sin páginas
MAX_FILAS_HOJA = 2000     # filas leídas por hoja de Excel / CSV

TEXTO = {".txt", ".md", ".json", ".py", ".sql", ".xml", ".yaml", ".yml", ".log", ".tsv"}
EXTENSIONES = {".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".csv", ".html", ".htm"} | TEXTO

# Formatos que Gemini puede leer directamente (sirve para PDF escaneados e imágenes).
MIME_GEMINI = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".txt": "text/plain",
    ".md": "text/plain",
    ".csv": "text/csv",
    ".html": "text/html",
    ".htm": "text/html",
}


class FormatoNoSoportado(ValueError):
    pass


@dataclass
class Seccion:
    numero: int
    etiqueta: str   # "página 3", "diapositiva 2", "hoja Backlog", "parte 1"
    texto: str


@dataclass
class Documento:
    nombre: str
    formato: str
    secciones: list[Seccion] = field(default_factory=list)
    aviso: str | None = None

    @property
    def caracteres(self) -> int:
        return sum(len(s.texto) for s in self.secciones)

    @property
    def tiene_texto(self) -> bool:
        return self.caracteres > 0


def extension(nombre: str) -> str:
    return PurePosixPath(nombre).suffix.lower()


def es_soportado(nombre: str) -> bool:
    return extension(nombre) in EXTENSIONES


def extraer(nombre: str, data: bytes) -> Documento:
    ext = extension(nombre)
    if ext == ".pdf":
        doc = _pdf(nombre, data)
    elif ext == ".docx":
        doc = _docx(nombre, data)
    elif ext == ".pptx":
        doc = _pptx(nombre, data)
    elif ext in (".xlsx", ".xlsm"):
        doc = _xlsx(nombre, data)
    elif ext in (".csv", ".tsv"):
        doc = _csv(nombre, data, "\t" if ext == ".tsv" else None)
    elif ext in (".html", ".htm"):
        doc = _partes(nombre, "html", _html_a_texto(_decodificar(data)))
    elif ext in TEXTO:
        doc = _partes(nombre, ext.lstrip("."), _decodificar(data))
    else:
        raise FormatoNoSoportado(
            f"No sé extraer texto de '{ext or nombre}'. Formatos: {', '.join(sorted(EXTENSIONES))}"
        )
    if not doc.tiene_texto and doc.aviso is None:
        doc.aviso = "El documento no tiene texto extraíble (¿escaneado o solo imágenes?)."
    return doc


# --------------------------------------------------------------------- formatos
def _pdf(nombre: str, data: bytes) -> Documento:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return Documento(nombre, "pdf", aviso="PDF protegido con contraseña.")
    secciones = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            texto = _limpiar(page.extract_text() or "")
        except Exception:
            texto = ""
        secciones.append(Seccion(i, f"página {i}", texto))
    doc = Documento(nombre, "pdf", secciones)
    vacias = sum(1 for s in secciones if not s.texto)
    if secciones and 0 < vacias < len(secciones):
        doc.aviso = f"{vacias} de {len(secciones)} páginas no tienen texto (posiblemente escaneadas)."
    return doc


def _docx(nombre: str, data: bytes) -> Documento:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = DocxDocument(io.BytesIO(data))
    lineas: list[str] = []
    for bloque in d.element.body.iterchildren():
        tag = bloque.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(bloque, d)
            if p.text.strip():
                estilo = (p.style.name or "").lower() if p.style is not None else ""
                prefijo = "## " if estilo.startswith(("heading", "título", "titulo", "title")) else ""
                lineas.append(prefijo + p.text.strip())
        elif tag == "tbl":
            for fila in Table(bloque, d).rows:
                celdas = [c.text.strip() for c in fila.cells]
                lineas.append(" | ".join(celdas))
    return _partes(nombre, "docx", "\n".join(lineas))


def _pptx(nombre: str, data: bytes) -> Documento:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    secciones = []
    for i, slide in enumerate(prs.slides, start=1):
        textos = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                textos.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for fila in shape.table.rows:
                    textos.append(" | ".join(c.text.strip() for c in fila.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notas = slide.notes_slide.notes_text_frame.text.strip()
            if notas:
                textos.append(f"[Notas] {notas}")
        secciones.append(Seccion(i, f"diapositiva {i}", _limpiar("\n".join(textos))))
    return Documento(nombre, "pptx", secciones)


def _xlsx(nombre: str, data: bytes) -> Documento:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    secciones = []
    aviso = None
    for i, ws in enumerate(wb.worksheets, start=1):
        filas = []
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= MAX_FILAS_HOJA:
                aviso = f"Algunas hojas se truncaron a {MAX_FILAS_HOJA} filas."
                break
            if row and any(v is not None and str(v).strip() for v in row):
                filas.append(" | ".join("" if v is None else str(v) for v in row))
        secciones.append(Seccion(i, f"hoja {ws.title}", "\n".join(filas)))
    wb.close()
    return Documento(nombre, "xlsx", secciones, aviso)


def _csv(nombre: str, data: bytes, delimitador: str | None) -> Documento:
    texto = _decodificar(data)
    if delimitador is None:
        try:
            delimitador = csv.Sniffer().sniff(texto[:5000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimitador = ","
    filas = []
    for n, row in enumerate(csv.reader(io.StringIO(texto), delimiter=delimitador)):
        if n >= MAX_FILAS_HOJA:
            break
        filas.append(" | ".join(row))
    doc = _partes(nombre, "csv", "\n".join(filas))
    if texto.count("\n") > MAX_FILAS_HOJA:
        doc.aviso = f"Solo se leyeron las primeras {MAX_FILAS_HOJA} filas."
    return doc


# ---------------------------------------------------------------------- ayudas
def _partes(nombre: str, formato: str, texto: str) -> Documento:
    """Divide texto sin páginas en partes de ~TAM_PARTE caracteres, cortando en saltos de línea."""
    texto = _limpiar(texto)
    secciones: list[Seccion] = []
    actual: list[str] = []
    largo = 0
    for linea in texto.split("\n"):
        if largo + len(linea) > TAM_PARTE and actual:
            secciones.append(Seccion(len(secciones) + 1, f"parte {len(secciones) + 1}", "\n".join(actual)))
            actual, largo = [], 0
        actual.append(linea)
        largo += len(linea) + 1
    if any(l.strip() for l in actual):
        secciones.append(Seccion(len(secciones) + 1, f"parte {len(secciones) + 1}", "\n".join(actual)))
    return Documento(nombre, formato, secciones)


def _decodificar(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _limpiar(texto: str) -> str:
    texto = texto.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


class _ExtractorHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.partes: list[str] = []
        self._omitir = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._omitir += 1
        elif tag in ("p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4"):
            self.partes.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._omitir:
            self._omitir -= 1

    def handle_data(self, data):
        if not self._omitir:
            self.partes.append(data)


def _html_a_texto(html: str) -> str:
    p = _ExtractorHTML()
    p.feed(html)
    return "".join(p.partes)
